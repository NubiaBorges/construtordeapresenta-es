#!/usr/bin/env python3
"""
Converte o roteiro.md em notas do apresentador e, opcionalmente, injeta no deck HTML.

Uso:
  python notas_do_roteiro.py roteiro.md --json notas.json
  python notas_do_roteiro.py roteiro.md --deck deck.html --saida deck-com-notas.html
  python notas_do_roteiro.py --listar-deck deck.html
  python notas_do_roteiro.py roteiro.md --pptx deck.pptx [--saida deck-com-notas.pptx]

O roteiro precisa seguir references/modelo-roteiro.md:
  ## Slide NN — Título · ⏱ tempo
  **Objetivo:** ...
  **Abordar:**
  - ...
  **Fala sugerida:** ...
  **Transição:** ...
"""
import argparse, html, json, re, sys

CAB = re.compile(r"^##\s+Slide\s+(\d+)\s*[—–-]\s*(.+?)\s*·\s*⏱\s*(.+?)\s*$")
ROTULO = re.compile(r"^\*\*(Objetivo|Abordar|Fala sugerida|Dados citados|Transição|Se perguntarem):\*\*\s*(.*)$")


def inline(txt):
    t = html.escape(txt.strip(), quote=False)
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"<i>\1</i>", t)
    return t


def ler_roteiro(caminho):
    slides, atual, campo = [], None, None
    for linha in open(caminho, encoding="utf-8").read().splitlines():
        m = CAB.match(linha)
        if m:
            atual = {"n": int(m.group(1)), "titulo": m.group(2).strip(), "tempo": m.group(3).strip(), "campos": {}}
            slides.append(atual); campo = None
            continue
        if linha.startswith("## ") or linha.startswith("# "):
            atual = None; campo = None  # fim dos blocos de slide (ex.: Preparação para perguntas)
            continue
        if atual is None:
            continue
        if linha.strip() == "---":
            campo = None; continue
        r = ROTULO.match(linha.strip())
        if r:
            campo = r.group(1)
            atual["campos"].setdefault(campo, [])
            if r.group(2).strip():
                atual["campos"][campo].append(r.group(2).strip())
            continue
        if campo and linha.strip():
            atual["campos"][campo].append(linha.strip())
    # checagens
    erros = []
    for i, s in enumerate(slides, 1):
        if s["n"] != i:
            erros.append(f"numeração fora de ordem: esperado Slide {i:02d}, encontrado Slide {s['n']:02d}")
        for obrig in ("Objetivo", "Abordar", "Fala sugerida"):
            if not s["campos"].get(obrig):
                erros.append(f"Slide {s['n']:02d}: falta **{obrig}:**")
    if not slides:
        erros.append("nenhum bloco '## Slide NN — Título · ⏱ tempo' encontrado")
    return slides, erros


def nota_html(s):
    c = s["campos"]
    partes = [f'<p class="tempo">⏱ {inline(s["tempo"])} · {inline(s["titulo"])}</p>']
    itens = [l[2:] if l.startswith("- ") else l for l in c.get("Abordar", [])]
    if itens:
        partes.append("<ul>" + "".join(f"<li>{inline(i)}</li>" for i in itens) + "</ul>")
    fala = " ".join(c.get("Fala sugerida", []))
    if fala:
        partes.append(f"<p>{inline(fala)}</p>")
    trans = " ".join(c.get("Transição", []))
    if trans:
        partes.append(f'<p style="opacity:.75"><b>→</b> {inline(trans)}</p>')
    return "".join(partes)


def nota_texto(s):
    """Versão em texto simples da nota, para as anotações do PowerPoint."""
    c = s["campos"]
    limpa = lambda x: re.sub(r"\*\*?(.+?)\*\*?", r"\1", x.strip())
    linhas = [f"⏱ {s['tempo']} · {s['titulo']}", ""]
    for l in c.get("Abordar", []):
        linhas.append("• " + limpa(l[2:] if l.startswith("- ") else l))
    fala = " ".join(c.get("Fala sugerida", []))
    if fala:
        linhas += ["", limpa(fala)]
    trans = " ".join(c.get("Transição", []))
    if trans:
        linhas += ["", "→ " + limpa(trans)]
    return "\n".join(linhas)


def notas_no_pptx(slides, caminho, destino):
    from pptx import Presentation
    prs = Presentation(caminho)
    total = len(prs.slides)
    if total != len(slides):
        lado = "o PowerPoint" if total > len(slides) else "o roteiro"
        sys.exit(f"DESALINHADO: PowerPoint tem {total} slides e roteiro tem {len(slides)} ({lado} tem slides a mais). Nada foi gravado.")
    for i, sl in enumerate(prs.slides):
        sl.notes_slide.notes_text_frame.text = nota_texto(slides[i])
    prs.save(destino)
    print(f"notas gravadas em {total} slides do PowerPoint → {destino}")
    for i, sl in enumerate(prs.slides, 1):
        def tamanho(sh):  # maior fonte usada na forma; o título costuma ser a maior
            tam = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
            if not tam and sh.is_placeholder and "TITLE" in str(sh.placeholder_format.type):
                return 999
            return max(tam) if tam else 0
        textos = [sh for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        titulo = max(textos, key=tamanho).text_frame.text if textos else "(sem texto)"
        print(f"  {i:02d}  pptx: {titulo.replace(chr(10), ' ')[:50]:50s} | roteiro: {slides[i-1]['titulo'][:50]}")


SEC_ABRE = re.compile(r'<section\b[^>]*class="[^"]*\bslide\b[^"]*"[^>]*>')


def slides_do_deck(doc):
    comentarios = [(m.start(), m.end()) for m in re.finditer(r"<!--.*?-->", doc, re.S)]
    dentro = lambda p: any(a <= p < b for a, b in comentarios)
    pos = []
    for m in SEC_ABRE.finditer(doc):
        if dentro(m.start()):
            continue
        fim = doc.find("</section>", m.end())
        if fim == -1:
            sys.exit("deck inválido: <section class=\"slide\"> sem </section>")
        pos.append((m.start(), fim))
    return pos


def titulo_slide(trecho):
    for tag in ("h1", "h2", "blockquote", "h3"):
        m = re.search(rf"<{tag}\b[^>]*>(.*?)</{tag}>", trecho, re.S)
        if m:
            return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", m.group(1))).strip()
    m = re.search(r'class="[^"]*stat-value[^"]*"[^>]*>(.*?)<', trecho, re.S)
    return m.group(1).strip() if m else "(sem título)"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("roteiro", nargs="?")
    ap.add_argument("--json")
    ap.add_argument("--deck")
    ap.add_argument("--saida")
    ap.add_argument("--listar-deck")
    ap.add_argument("--pptx")
    a = ap.parse_args()

    if a.listar_deck:
        doc = open(a.listar_deck, encoding="utf-8").read()
        for i, (ini, fim) in enumerate(slides_do_deck(doc), 1):
            print(f"{i:02d} — {titulo_slide(doc[ini:fim])}")
        return

    if not a.roteiro:
        ap.error("informe o roteiro.md")
    slides, erros = ler_roteiro(a.roteiro)
    if erros:
        print("Roteiro fora do modelo:"); [print("  -", e) for e in erros]; sys.exit(1)
    notas = {s["n"]: nota_html(s) for s in slides}
    print(f"{len(slides)} slides lidos do roteiro.")

    if a.json:
        json.dump({str(k): v for k, v in notas.items()}, open(a.json, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        print(f"notas salvas em {a.json}")

    if a.pptx:
        notas_no_pptx(slides, a.pptx, a.saida if (a.saida and not a.deck) else a.pptx)

    if a.deck:
        doc = open(a.deck, encoding="utf-8").read()
        if 'id="notesPanel"' not in doc:
            print("AVISO: este deck não tem o painel de notas (tecla N). As notas serão injetadas, mas só aparecem "
                  "se o deck for reconstruído com o template atual da construtor-apresentacoes.")
        pos = slides_do_deck(doc)
        if len(pos) != len(slides):
            lado = "o deck" if len(pos) > len(slides) else "o roteiro"
            sys.exit(f"DESALINHADO: deck tem {len(pos)} slides e roteiro tem {len(slides)} ({lado} tem slides a mais). Nada foi injetado.")
        saida, ultimo = [], 0
        for i, (ini, fim) in enumerate(pos, 1):
            corpo = re.sub(r'\s*<aside class="notes">.*?</aside>', "", doc[ultimo:fim], flags=re.S)
            saida.append(corpo + f'\n      <aside class="notes">{notas[i]}</aside>\n    ')
            ultimo = fim
        saida.append(doc[ultimo:])
        destino = a.saida or a.deck
        open(destino, "w", encoding="utf-8").write("".join(saida))
        print(f"notas injetadas em {len(pos)} slides → {destino}")
        print("Confira títulos lado a lado:")
        for i, (ini, fim) in enumerate(pos, 1):
            print(f"  {i:02d}  deck: {titulo_slide(doc[ini:fim])[:50]:50s} | roteiro: {slides[i-1]['titulo'][:50]}")


if __name__ == "__main__":
    main()
